# ==========================================
# OM Automation V2
# api/gst_captcha.py
#
# Captcha-driven GST verification against the real portal (services.gst.gov.in) via gst_search.py.
# One real (headless) browser per user, one GSTIN at a time: show captcha -> user types it -> submit ->
# save result -> show next captcha. The captcha is never auto-solved.
#
# Web-app changes vs. the original: sessions are keyed by the logged-in user (not a Flask cookie),
# the number of simultaneous Chrome sessions is capped (RAM) and idle sessions are closed automatically.
# State lives in this process (the app runs one gunicorn worker).
# ==========================================

import hashlib
import os
import re
import threading
import time

from openpyxl import Workbook, load_workbook

from omauto import config
from omauto.config import GST_OUTPUT_COLUMNS, GSTIN_HEADERS
from omauto.excel.generic_reader import GenericReader
from omauto.gst_search import (
    open_browser, go_to_search_page, fill_gst,
    screenshot_captcha, fill_captcha, click_search,
    get_page_text, looks_like_wrong_captcha, extract_result,
    close_browser,
)

SESSIONS = {}
LOCK = threading.Lock()
MAX_SESSIONS = int(os.environ.get("GST_MAX_SESSIONS", "1"))      # each headless Chrome needs ~300 MB
IDLE_TIMEOUT = int(os.environ.get("GST_IDLE_SECONDS", "600"))

_GSTIN_RE = re.compile(r"^[0-9A-Za-z]{15}$")


def is_valid_gstin(value):
    """Structural check only (15 alphanumeric characters) - not a full checksum validation."""
    return bool(_GSTIN_RE.match((value or "").strip()))


def _sid(uid):
    return hashlib.sha1(str(uid).encode()).hexdigest()[:16]


def get_state(uid):
    return SESSIONS.get(_sid(uid))


def output_path(uid):
    return os.path.join(config.OUTPUT_DIR, f"GST_Report_{_sid(uid)}.xlsx")


def report_path(uid):
    p = output_path(uid)
    return p if os.path.exists(p) else None


def _captcha_path(uid):
    return os.path.join(config.CAPTCHA_DIR, f"{_sid(uid)}.png")


def captcha_image_path(uid):
    p = _captcha_path(uid)
    return p if os.path.exists(p) and get_state(uid) else None


def _flatten(value):
    """Collapse line breaks so every row of the downloaded report stays one line tall."""
    if isinstance(value, str):
        return value.replace("\r\n", " | ").replace("\n", " | ").replace("\r", " | ")
    return value


def append_result(path, result):
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(GST_OUTPUT_COLUMNS)
    ws.append([_flatten(result.get(col, "")) for col in GST_OUTPUT_COLUMNS])
    wb.save(path)


def cleanup_stale():
    """Close browsers nobody has touched for IDLE_TIMEOUT seconds (frees RAM for the next user)."""
    now = time.time()
    for sid, st in list(SESSIONS.items()):
        if now - st.get("last", now) > IDLE_TIMEOUT:
            if st.get("driver"):
                close_browser(st["driver"])
            SESSIONS.pop(sid, None)


def _advance(state, uid):
    """Move to the next GSTIN: fill it in and get a fresh captcha ready."""
    state["current_index"] += 1
    idx, total = state["current_index"], len(state["gst_list"])
    if idx >= total:
        close_browser(state["driver"])
        state["driver"] = None
        return {"status": "done", "total": total, "processed": total,
                "output_filename": os.path.basename(state["output_path"])}
    gst, driver = state["gst_list"][idx], state["driver"]
    go_to_search_page(driver)
    fill_gst(driver, gst)
    screenshot_captcha(driver, _captcha_path(uid))
    return {"status": "waiting_captcha", "gst": gst, "index": idx, "total": total}


def start_session(uid, gst_list):
    """gst_list: list of already-cleaned GSTIN strings."""
    with LOCK:
        cleanup_stale()
        old = SESSIONS.pop(_sid(uid), None)           # a user only ever has one session: replace it
        if old and old.get("driver"):
            close_browser(old["driver"])
        busy = sum(1 for s in SESSIONS.values() if s.get("driver"))
        if busy >= MAX_SESSIONS:
            raise RuntimeError("GST browser abhi doosre user ke use me hai (server par ek saath limited sessions). "
                               "Kuch minute baad dobara try karein.")
        # a fresh report file per session (the original behaviour)
        try:
            os.remove(output_path(uid))
        except OSError:
            pass
        driver = open_browser()
        state = {"gst_list": gst_list, "current_index": -1, "driver": driver,
                 "output_path": output_path(uid), "last": time.time()}
        SESSIONS[_sid(uid)] = state
        try:
            return _advance(state, uid)
        except Exception:
            close_browser(driver)
            SESSIONS.pop(_sid(uid), None)
            raise


def start_excel(uid, file_storage):
    records = GenericReader(file_storage, GSTIN_HEADERS, label="GSTIN").read()
    gst_list, seen = [], set()
    for r in records:
        value = r["value"].strip().upper()
        if value in seen:
            continue
        seen.add(value)
        gst_list.append(value)
    if not gst_list:
        return {"error": "No GSTIN column found in that file."}
    return start_session(uid, gst_list)


def submit_captcha(uid, captcha_text):
    state = get_state(uid)
    if not state or not state.get("driver"):
        return {"error": "Session expired. Please start again."}, 400
    captcha_text = (captcha_text or "").strip()
    if not captcha_text:
        return {"error": "Please type the captcha text."}, 400

    with LOCK:
        state["last"] = time.time()
        driver = state["driver"]
        gst = state["gst_list"][state["current_index"]]
        try:
            fill_captcha(driver, captcha_text)
            click_search(driver)
        except Exception as e:  # noqa: BLE001
            return {"error": f"Could not submit captcha: {e}"}, 500

        page_text = get_page_text(driver)
        if looks_like_wrong_captcha(page_text):
            # Same GSTIN, fresh captcha - don't advance the index.
            go_to_search_page(driver)
            fill_gst(driver, gst)
            screenshot_captcha(driver, _captcha_path(uid))
            return {"status": "retry_captcha", "message": "Captcha didn't match - a new one is ready, try again.",
                    "gst": gst, "index": state["current_index"], "total": len(state["gst_list"])}, 200

        result = extract_result(driver, gst)
        append_result(state["output_path"], result)
        payload = _advance(state, uid)
        payload["last_result"] = {"gst": gst, "remarks": result.get("Remarks", "")}
        if payload["status"] != "done":
            payload["processed"] = state["current_index"]
        return payload, 200


def stop_session(uid):
    state = get_state(uid)
    if not state:
        return {"error": "Session expired."}, 400
    with LOCK:
        if state.get("driver"):
            close_browser(state["driver"])
            state["driver"] = None
        processed = max(state["current_index"], 0)
        total = len(state["gst_list"])
        path = state["output_path"]
        return {"ok": True, "processed": processed, "total": total,
                "output_filename": os.path.basename(path) if os.path.exists(path) else None}, 200
