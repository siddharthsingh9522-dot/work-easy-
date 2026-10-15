"""Subject / body / attachment builders of the three mail workflows (ported from om-suite's main_app/app.py)."""
import csv
import io
import re

from . import bill_parser, previous_mail

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TBB_INTRO = (
    "Dear sir,\n\n"
    "Please check the GR. If this GR belongs to a TBB customer, kindly provide the "
    "correct TBB Party Code and modify the same accordingly.\n\n"
    "Please confirm once updated.\n\n"
)


def _esc(val):
    return str(val or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def parse_recipients(value):
    """Comma / semicolon / newline separated addresses -> lower-cased, de-duplicated list. Invalid address -> ValueError."""
    out, seen = [], set()
    for item in re.split(r"[,;\n]+", str(value or "")):
        email = item.strip().lower()
        if not email:
            continue
        if not _EMAIL_RE.match(email):
            raise ValueError(f"Invalid email address: {email}")
        if email not in seen:
            seen.add(email)
            out.append(email)
    return out


# ------------------------------------------------------------------ TBB party mail
def tbb_subject(branch):
    grs = branch.get("grs", [])
    cns = [str(gr.get("cn", "")).strip() for gr in grs if str(gr.get("cn", "")).strip()]
    gr_part = f"{cns[0]}/{cns[-1]}" if len(cns) >= 2 else (cns[0] if cns else "")
    names = []
    for gr in grs:
        pn = str(gr.get("party_name", "")).strip()
        if pn and pn.lower() != "nan" and pn not in names:
            names.append(pn)
    return f"PROVIDE THE TBB CODE & MODIFY THE SAME............. GR:-{gr_part} M/S- {' & '.join(names)}"


def tbb_text(branch, from_name=""):
    body = TBB_INTRO + "B.Code \tBRANCH \tCN_NO \tPARTY CODE \tPARTY NAME\n"
    for gr in branch["grs"]:
        body += (f"{branch['branch_code']} \t{branch.get('branch_name_sheet', '')} \t"
                 f"{gr['cn']} \t{gr['party_code']} \t*{gr['party_name']}*\n")
    return body + "\nRegards,\n" + (from_name or "")


def tbb_html(branch, from_name=""):
    intro = "".join(f"<p>{_esc(line)}</p>" for line in TBB_INTRO.strip("\n").split("\n\n"))
    th = "background:#1F4E78;color:#fff;padding:6px 10px;font-size:12px;border:1px solid #333;text-align:left;"
    td = "padding:6px 10px;font-size:12px;border:1px solid #333;"
    head = "".join(f'<th style="{th}">{_esc(h)}</th>' for h in ["B.Code", "BRANCH", "CN_NO", "PARTY CODE", "PARTY NAME"])
    rows = []
    for i, gr in enumerate(branch["grs"]):
        bg = "#F2F2F2" if i % 2 == 0 else "#FFFFFF"
        cell = f'style="{td}background:{bg};"'
        party = f'style="{td}background:{bg};color:#C00000;font-weight:bold;"'
        rows.append(f"<tr><td {cell}>{_esc(branch['branch_code'])}</td><td {cell}>{_esc(branch.get('branch_name_sheet', ''))}</td>"
                    f"<td {cell}>{_esc(gr['cn'])}</td><td {cell}>{_esc(gr['party_code'])}</td><td {party}>{_esc(gr['party_name'])}</td></tr>")
    table = (f'<table style="border-collapse:collapse;border:1px solid #333;margin-top:6px;"><thead><tr>{head}</tr></thead>'
             f"<tbody>{''.join(rows)}</tbody></table>")
    return (f"<html><body style='font-family:Arial,sans-serif;font-size:13px;'>{intro}{table}"
            f"<p>Regards,<br>{_esc(from_name)}</p></body></html>")


def tbb_xlsx(branch):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "TBB Details"
    headers = ["Branch Code", "Branch Name", "Region", "CN No", "Party Code", "Party Name", "CN Date"]
    ws.append(headers)
    for gr in branch["grs"]:
        ws.append([branch["branch_code"], branch.get("branch_name_sheet", ""), branch.get("region", ""),
                   gr["cn"], gr["party_code"], gr["party_name"], gr["cn_date"]])
    for i, h in enumerate(headers, start=1):
        width = max([len(h)] + [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width + 4
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(columns, rows):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


# ------------------------------------------------------------------ one mail dict per workflow row
def build_mail(module, code, branch, from_name=""):
    """-> dict understood by dbp.mail_sender.build_message and core.bridge_queue.enqueue_send (recipients added later)."""
    if module == "tbb":
        name = f"TBB_{code}_{branch.get('branch_name_sheet', '')}".replace(" ", "_") + ".xlsx"
        return dict(subject=tbb_subject(branch), body=tbb_text(branch, from_name), html_body=tbb_html(branch, from_name),
                    attachment_bytes=tbb_xlsx(branch), attachment_name=name, attachment_type=XLSX)
    if module == "bill":
        att = bill_parser.build_attachment_csv(branch)
        return dict(subject=bill_parser.build_subject(branch), body=bill_parser.build_email_text(branch),
                    html_body=bill_parser.build_email_html(branch), attachment_bytes=att,
                    attachment_name=f"Bill_{code}.csv", attachment_type="text/csv")
    if module == "previous":
        cols = branch.get("columns", [])
        att = _csv(cols, [[r.get(c, "") for c in cols] for r in branch.get("rows", [])]) if branch.get("rows") else None
        return dict(subject=branch.get("subject", ""), body=branch.get("body", ""), html_body="",
                    attachment_bytes=att, attachment_name=f"PreviousMail_{branch.get('key', code)}".replace(" ", "_") + ".csv",
                    attachment_type="text/csv")
    raise ValueError(module)


def fingerprint_text(module, branch, from_name=""):
    """What the idempotency fingerprint covers (changing the visible content makes a fresh DRAFT)."""
    if module == "tbb":
        return tbb_text(branch, from_name)
    if module == "bill":
        return bill_parser.build_email_text(branch)
    return branch.get("subject", "") + branch.get("body", "")
