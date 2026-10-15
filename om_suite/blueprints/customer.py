"""Customer Search - manual search (name / codes / GST) and bulk Excel by booking-location code."""
import os
import re
import time

import openpyxl
from urllib.parse import urlsplit
import requests
from bs4 import BeautifulSoup
from flask import Blueprint, current_app, g, jsonify, render_template, request
from openpyxl.utils import get_column_letter

from core import jobs
from core.auth import module_required

bp = Blueprint("customer", __name__, url_prefix="/customer")

DELAY_SECONDS, TIMEOUT_SECONDS, NOT_FOUND_WAIT = 0.4, 30, 15
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Content-Type": "application/x-www-form-urlencoded",
}
RESULT_COLUMNS = [
    "CCode", "CNAME", "Address", "Booking Location", "Acc Loc", "Bill Loc", "Verify", "Type", "GST NO", "Status",
    "Master Code", "Kom", "Document", "First Booking date", "Vendor Code", "CUSTOMER MOBILE NO", "MERGE CODE",
    "DELIVERY CODE", "RETAIL TRANSFER DATE", "CONT RATE", "CUSTOMER ENTER BY", "CUSTOMER ENTER DATE",
]
_ILLEGAL = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")
MAX_CODES = 5000


@bp.before_request
def _active():
    g.active_module = "customer_search"


def clean(value):
    return _ILLEGAL.sub("", value) if isinstance(value, str) else value


def search_customer(session, url, cname="", ccode="", lcode="", blcode="", mcode="", custgst=""):
    payload = {"cname": cname, "ccode": ccode, "lcode": lcode, "blcode": blcode, "mcode": mcode,
               "custgst": custgst, "submit": "Search"}
    parts = urlsplit(url)
    headers = {**HEADERS, "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
               "Accept-Language": "en-GB,en-US;q=0.9,en;q=0.8", "Origin": f"{parts.scheme}://{parts.netloc}", "Referer": url}
    resp = session.post(url, data=payload, headers=headers, timeout=TIMEOUT_SECONDS)
    resp.raise_for_status()
    no_data = "no data found" in resp.text.lower()
    table = BeautifulSoup(resp.text, "html.parser").find("table")
    if table is None:
        return [], no_data
    rows = []
    for tr in (table.find("tbody") or table).find_all("tr"):
        cells = tr.find_all("td")
        if not cells:
            continue
        vals = [c.get_text(strip=True) for c in cells]
        vals = (vals + [""] * len(RESULT_COLUMNS))[:len(RESULT_COLUMNS)]
        rows.append(dict(zip(RESULT_COLUMNS, vals)))
    return rows, no_data


def search_with_retry(session, url, code, job=None):
    """Original behaviour: an empty answer that is not an explicit 'no data found' is retried once after a pause."""
    try:
        rows, no_data = search_customer(session, url, blcode=code)
    except requests.RequestException as e:
        current_app.logger.warning("customer search(%s) failed: %s: %s", code, type(e).__name__, e)
        rows, no_data = [], False
    if rows or no_data:
        return rows
    for _ in range(NOT_FOUND_WAIT):
        if job is not None and job.stopped:
            return rows
        time.sleep(1)
    try:
        rows, _ = search_customer(session, url, blcode=code)
    except requests.RequestException:
        rows = []
    return rows


def find_code_column(ws):
    keywords = ["ccode", "customer code", "cust code", "code", "customer_code"]
    for idx, header in enumerate([c.value for c in ws[1]], start=1):
        if header and str(header).strip().lower() in keywords:
            return idx, True
    best_col, best = None, 0
    for col in range(1, ws.max_column + 1):
        n = sum(1 for r in range(1, min(ws.max_row, 50) + 1)
                if ws.cell(row=r, column=col).value is not None
                and str(ws.cell(row=r, column=col).value).strip().replace(".", "", 1).isdigit())
        if n > best:
            best, best_col = n, col
    if best_col is None:
        raise ValueError("Could not detect a column containing customer codes.")
    first = ws.cell(row=1, column=best_col).value
    return best_col, not (first is not None and str(first).strip().replace(".", "", 1).isdigit())


def read_customer_codes(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    col, has_header = find_code_column(ws)
    codes = []
    for r in range(2 if has_header else 1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if v is None or str(v).strip() == "":
            continue
        s = str(v).strip()
        codes.append(s[:-2] if s.endswith(".0") else s)
    return codes


def write_results_excel(path, all_results):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Searched Code"] + RESULT_COLUMNS)
    for c in range(1, len(RESULT_COLUMNS) + 2):
        ws.cell(row=1, column=c).font = openpyxl.styles.Font(bold=True)
    not_found = []
    for code, rows in all_results:
        code = clean(code)
        if not rows:
            ws.append([code, "NOT FOUND"] + [""] * (len(RESULT_COLUMNS) - 1))
            not_found.append(code)
            continue
        for row in rows:
            ws.append([code] + [clean(row[c]) for c in RESULT_COLUMNS])
    for c in range(1, len(RESULT_COLUMNS) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 18
    if not_found:
        ws2 = wb.create_sheet("Not Found")
        ws2.append(["Customer Code"])
        for c in not_found:
            ws2.append([c])
    wb.save(path)


@bp.route("/")
@module_required("customer_search")
def index():
    return render_template("customer/index.html", columns=RESULT_COLUMNS)


@bp.route("/api/manual-search", methods=["POST"])
@module_required("customer_search")
def manual_search():
    d = request.get_json(silent=True) or {}
    fields = {k: str(d.get(k, ""))[:120] for k in ("cname", "ccode", "lcode", "blcode", "mcode", "custgst")}
    if not any(v.strip() for v in fields.values()):
        return jsonify(ok=False, error="Kam se kam ek search field bharein."), 400
    try:
        rows, _ = search_customer(requests.Session(), current_app.config["CUSTOMER_URL"], **fields)
        return jsonify(ok=True, rows=rows, columns=RESULT_COLUMNS)
    except requests.RequestException as e:
        return jsonify(ok=False, error=f"Portal se jawab nahi mila: {e}"), 502


def _bulk(job, path, url):
    codes = read_customer_codes(path)[:MAX_CODES]
    os.remove(path)
    if not codes:
        raise ValueError("No customer codes found in the file")
    job.progress(0, len(codes))
    session, results = requests.Session(), []
    for code in codes:
        if job.stopped:
            break
        rows = search_with_retry(session, url, code, job)
        results.append((code, rows))
        job.bump("found" if rows else "not_found")
        job.progress(job.current + 1)
        time.sleep(DELAY_SECONDS)
    write_results_excel(job.path("customer_search_results.xlsx"), results)
    job.output_name = "customer_search_results.xlsx"


@bp.route("/api/bulk-search/start", methods=["POST"])
@module_required("customer_search")
def bulk_start():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(ok=False, error="No file selected"), 400
    if os.path.splitext(f.filename)[1].lower() not in (".xlsx", ".xlsm"):
        return jsonify(ok=False, error="Sirf Excel (.xlsx) file."), 400
    path = jobs.save_upload(f, ".xlsx")
    url = current_app.config["CUSTOMER_URL"]
    job = jobs.start(g.user["email"], "customer-bulk", lambda j: _bulk(j, path, url))
    return jsonify(ok=True, job_id=job.id)
