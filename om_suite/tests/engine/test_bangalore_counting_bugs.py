"""
Regression tests for the 3 bugs reported on 2026-10-10 (BANGALORE export):
 1. NET data not read properly -- sheet named "NET POD" etc. was treated as a
    party sheet, and CN_DELIVERY_DATE like '14-JUL-26 03.23.58.000000 AM'
    could not be parsed (whole NET fell into the "current month").
 2. Repeated header rows / 'BRANCH: 502 - JIGNI' marker rows inside a sheet
    were counted as data -> fake 'ADR BRANCH' / 'BILL_BRANCH_CODE' branches.
 3. Summary therefore wrong (extra pending, wrong months, missing zero rows).
"""
import openpyxl
import pandas as pd
import pytest

from dbp import ingest
from dbp import month_summary as ms
from dbp import region_branches as rb
# ---------------------------------------------------------------- sheet names
@pytest.mark.parametrize("name", ["NET", "NET POD", "NETPOD", "POD NET", "NET-POD",
                                  "NET (POD)", "POD", "net pod", "Net Pod\n",
                                  "NET POD MADE BUT NOT BILLED", "POD ON NET BUT NOT BILLED"])
def test_net_sheet_name_variants(name):
    assert ingest.categorize_sheet(name) == "POD_NOT_BILLED"


@pytest.mark.parametrize("name", ["ADR", "ADR MADE", "adr", "ADR MADE BUT NOT BILLED", "PENDING BILLING"])
def test_adr_sheet_name_variants(name):
    assert ingest.categorize_sheet(name) == "PENDING_BILLING"


@pytest.mark.parametrize("name", ["NET ONE TECHNOLOGIES", "TATA MOTOR", "BOSCH", "SUPRAJIT ", "ADR NET"])
def test_party_names_not_mistaken_for_category(name):
    assert ingest.categorize_sheet(name) == "OTHER"


# ---------------------------------------------------------------- date parsing
@pytest.mark.parametrize("raw,expected", [
    ("14-JUL-26 03.23.58.000000 AM", "2026-07-14"),
    ("DELIVERED-ON-15-JUN-26 12:56", "2026-06-15"),
    ("16-JUN-26", "2026-06-16"),
    ("21-AUGUST-2026", "2026-08-21"),
    ("2026-08-21 00:00:00", "2026-08-21"),
    ("06.06.2026", "2026-06-06"),
    ("10.07.2026", "2026-07-10"),
    ("13/07/26", "2026-07-13"),
    (46223, "2026-07-20"),
])
def test_parse_date_value_formats(raw, expected):
    assert ms.parse_date_value(raw) == pd.Timestamp(expected)


@pytest.mark.parametrize("raw", [None, "", "nan", "ADR DATE", "3401", "garbage", "99-XXX-26"])
def test_parse_date_value_rejects_non_dates(raw):
    assert pd.isna(ms.parse_date_value(raw))


# ---------------------------------------------------------------- structural rows
def _multi_block_sheet(wb, title, header, rows_per_block, n_blocks=3, branch_col_idx=3):
    ws = wb.create_sheet(title)
    for b in range(n_blocks):
        code = 500 + b
        ws.append([f"BRANCH: {code} - NAME{b}"] + [None] * (len(header) - 1))
        ws.append(header)
        for i in range(rows_per_block):
            row = [code, f"CN{b}{i}", "x", code] if len(header) == 4 else [code, f"CN{b}{i}"]
            ws.append(row)
    return ws


def test_repeated_header_and_marker_rows_never_counted(tmp_path):
    path = str(tmp_path / "multi.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _multi_block_sheet(wb, "ADR", ["BRANCH CODE", "CN_NO", "ADR DATE", "ADR BRANCH"], 4)
    wb.save(path)
    data = ingest.load_workbook_categorized(path)
    pb = data["PENDING_BILLING"]
    assert len(pb) == 12                               # 3 blocks x 4 rows, nothing else
    assert "ADR BRANCH" not in set(pb["ADR BRANCH"].astype(str))
    assert not pb["BRANCH CODE"].astype(str).str.upper().str.startswith("BRANCH").any()


def test_bogus_branch_text_not_added_as_branch_in_month_summary():
    df = pd.DataFrame({"ADR BRANCH": ["503", "ADR BRANCH", "", "BILL_BRANCH_CODE", "503"],
                       "ADR DATE": ["10-JUL-26", "ADR DATE", "", "x", "11-JUL-26"]})
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE"], [], current_month=pd.Period("2026-07", freq="M"))
    codes = list(mat["B.CODE"])
    assert codes == ["503", "TOTAL"]
    assert int(mat.iloc[0]["TOTAL"]) == 2
    audit = mat.attrs["audit"]
    assert audit["rows_dropped_bad_branch"] == 3
    assert "ADR BRANCH" in audit["dropped_samples"]


# ---------------------------------------------------------------- NET end to end
def test_net_oracle_timestamp_goes_to_right_month_not_current(tmp_path):
    path = str(tmp_path / "net.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NET POD"
    ws.append(["S.NO", "CN_NO", "CN_DELIVERY_DATE", "BILL_BRANCH_CODE"])
    ws.append([1, "N1", "14-JUL-26 03.23.58.000000 AM", 503])
    ws.append([2, "N2", "02-JUN-26 11.00.00.000000 PM", 503])
    ws.append([3, "N3", None, 503])
    wb.save(path)
    data = ingest.load_workbook_categorized(path)
    assert len(data["POD_NOT_BILLED"]) == 3
    mat, n_missing, labels = ms.build_net_month_summary(
        data, current_month=pd.Period("2026-08", freq="M"))
    assert n_missing == 1                              # only the truly date-less row
    row = mat[mat["B.CODE"] == "503"].iloc[0]
    assert row["June - 26"] == 1 and row["July - 26"] == 1 and row["August - 26"] == 1


def test_net_never_uses_adr_combo_date_as_fallback():
    df = pd.DataFrame({"BILL_BRANCH_CODE": [503],
                       "CN_DELIVERY_DATE": [None],
                       "ADR_NO_DATE_BRANCH": ["1,05-JAN-26,503"]})
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "BILL_BRANCH_CODE", ["CN_DELIVERY_DATE"], [], current_month=pd.Period("2026-08", freq="M"))
    assert n_missing == 1 and labels == ["August - 26"]


# ---------------------------------------------------------------- region branches
def test_region_branches_zero_rows_in_given_order_and_numeric_sort():
    df = pd.DataFrame({"ADR BRANCH": ["1802", "503", "9999"],
                       "ADR DATE": ["10-AUG-26"] * 3})
    mat, _, _ = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE"], [], region_branches=["501", "503", "1802"])
    assert list(mat["B.CODE"]) == ["501", "503", "1802", "9999", "TOTAL"]   # extra data branch kept
    assert int(mat[mat["B.CODE"] == "501"].iloc[0]["TOTAL"]) == 0
    # without a list: numeric (not string) order
    mat2, _, _ = ms.build_month_wise_summary(df, "ADR BRANCH", ["ADR DATE"], [])
    assert list(mat2["B.CODE"]) == ["503", "1802", "9999", "TOTAL"]


def test_region_branches_store_and_parsers(tmp_path, monkeypatch):
    monkeypatch.setenv("DBP_MASTER_DIR", str(tmp_path))
    assert rb.parse_codes_text("501, 502\n503;504 abc 504 0") == ["501", "502", "503", "504"]
    assert rb.save("Bangalore  region", ["501", "502", "x", "502"]) is True
    assert rb.get("BANGALORE REGION") == ["501", "502"]
    rb.delete("bangalore region")
    assert rb.get("BANGALORE REGION") == []

    p = str(tmp_path / "prev.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["title"])
    ws.append(["S No.", "B.CODE", "BRANCH NAME", "April - 25"])
    ws.append(["*", "TATA MOTOR", None, 1])
    ws.append([1, 501, "A", 0])
    ws.append([2, 502, "B", 0])
    ws.append(["TOTAL", None, None, 0])
    wb.save(p)
    assert rb.codes_from_excel(p) == ["501", "502"]


def test_full_report_audit_is_consistent(tmp_path):
    path = str(tmp_path / "full.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _multi_block_sheet(wb, "ADR", ["BRANCH CODE", "CN_NO", "ADR DATE", "ADR BRANCH"], 5)
    ws = wb.create_sheet("BOSCH")
    ws.append([501, "ADR MADE BUT NOT BILLED"])
    ws.append(["S.NO", "CN_NO", "ADR DATE", "BRANCH CODE"])
    for i in range(3):
        ws.append([i + 1, f"B{i}", "10-JUL-26", 501])
    wb.save(path)
    data = ingest.load_workbook_categorized(path)
    rep = ms.build_full_month_wise_report(data, {}, pd.Period("2026-07", freq="M"),
                                          pd.Period("2026-07", freq="M"))
    df, _, _ = rep["PENDING_BILLING"]
    a = df.attrs["audit"]
    assert a["rows_counted"] == 15 and a["party_rows_counted"] == 3
    assert a["grand_total"] == a["rows_counted"] + a["party_rows_counted"] == 18
