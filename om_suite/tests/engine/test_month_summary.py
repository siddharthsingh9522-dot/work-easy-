import pandas as pd

from dbp import month_summary as ms


def test_happy_path_explicit_date_column():
    df = pd.DataFrame({
        "ADR BRANCH": [101, 101, 102],
        "ADR DATE": ["21-AUG-26", "08-SEP-26", "15-JUL-26"],
    })
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE"], ["ADR_NO_DATE_BRANCH"],
        branch_name_map={101: "X", 102: "Y"}, current_month=pd.Period("2026-10", freq="M"))
    assert n_missing == 0
    row101 = mat[mat["B.CODE"] == "101"].iloc[0]
    assert row101["TOTAL"] == 2
    assert row101["August - 26"] == 1 and row101["September - 26"] == 1
    total_row = mat[mat["B.CODE"] == "TOTAL"].iloc[0]
    assert total_row["TOTAL"] == 3


def test_combo_column_fallback_when_no_explicit_date_col():
    df = pd.DataFrame({
        "ADR BRANCH": [530, 530],
        "ADR_NO_DATE_BRANCH": ["19202517000399,29-JAN-26,530", "19202517000400,15-FEB-26,530"],
    })
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE"], ["ADR_NO_DATE_BRANCH"],
        current_month=pd.Period("2026-10", freq="M"))
    assert n_missing == 0
    assert "January - 26" in labels and "February - 26" in labels


def test_rows_with_no_date_anywhere_go_to_current_month():
    df = pd.DataFrame({
        "ADR BRANCH": [201, 201],
        "ADR DATE": ["05-JUN-26", None],
        "ADR_NO_DATE_BRANCH": [None, None],
    })
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE"], ["ADR_NO_DATE_BRANCH"],
        current_month=pd.Period("2026-10", freq="M"))
    assert n_missing == 1
    row = mat[mat["B.CODE"] == "201"].iloc[0]
    assert row["June - 26"] == 1
    assert row["October - 26"] == 1
    assert row["TOTAL"] == 2


def test_mislabeled_date_column_is_detected_and_rejected():
    """Real bug (2026-10-02, JAMSHEDPUR export): a column literally named
    'DATE' held branch codes, while a column named 'ADR' held the real
    dates. The name-based alias match must not be trusted blindly -- it
    should validate the column actually looks like dates, and fall back to
    the combo-column when it doesn't."""
    df = pd.DataFrame({
        "ADR BRANCH": [3401, 3401],
        "DATE": ["3401", "3401"],                      # misleading column name
        "ADR_NO_DATE_BRANCH": ["1,21-AUG-26,3401", "2,19-AUG-26,3401"],
    })
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "ADR BRANCH", ["ADR DATE", "DATE"], ["ADR_NO_DATE_BRANCH"],
        current_month=pd.Period("2026-08", freq="M"))
    assert n_missing == 0
    assert labels == ["August - 26"]
    row = mat[mat["B.CODE"] == "3401"].iloc[0]
    assert row["August - 26"] == 2


def test_build_adr_month_summary_reuses_branch_resolution_and_exclusion(tmp_path):
    """End-to-end through ingest.load_workbook_categorized() -- confirms
    month_summary automatically benefits from the ADR BRANCH resolution and
    the OTHER END/DISPUTE MATTER exclusion fix, no duplicate logic."""
    import openpyxl
    from dbp import ingest

    path = str(tmp_path / "repro.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PENDING BILLING"
    rows = [[101, "ADR MADE BUT NOT BILLED"] + [None] * 5, [None] * 7,
            ["S.NO", "CN_NO", "ADR DATE", "ADR BRANCH", None, None, None]]
    for i in range(1, 4):
        rows.append([i, f"CN{i}", "10-AUG-26", 101, None, None, None])
    rows += [[None] * 7, ["DISPUTE MATTER"] + [None] * 6, [None] * 7,
             ["S.NO", "CN_NO", "ADR DATE", "ADR BRANCH", None, None, None]]
    for i in range(1, 3):
        rows.append([i, f"DISP{i}", "11-AUG-26", 101, None, None, None])
    for r in rows:
        ws.append(r)
    wb.save(path)

    categorized = ingest.load_workbook_categorized(path)
    mat, n_missing, labels = ms_mod_build(categorized)
    row = mat[mat["B.CODE"] == "101"].iloc[0]
    assert row["TOTAL"] == 3   # dispute-matter rows excluded, not counted


def ms_mod_build(categorized):
    import pandas as pd
    return ms.build_adr_month_summary(categorized, current_month=pd.Period("2026-10", freq="M"))


def test_net_uses_cn_delivery_date_and_bill_branch_code():
    """CONFIRMED by user (2026-10-03): NET month-wise summary CN_DELIVERY_DATE
    se date leta hai aur BILL_BRANCH_CODE se branch (POD_NOT_BILLED ka normal
    branch-of-truth column, ingest.detect_branch_column se hi aata hai)."""
    df = pd.DataFrame({
        "BILL_BRANCH_CODE": [3402, 3402, 3405],
        "CUSTOMER_BILL_LOC": [9999, 9999, 9999],  # decoy -- NOT the branch column for NET
        "CN_DELIVERY_DATE": ["17-AUG-26", None, "15-SEP-25"],
    })
    mat, n_missing, labels = ms.build_month_wise_summary(
        df, "BILL_BRANCH_CODE", ["CN_DELIVERY_DATE"], [],
        current_month=pd.Period("2026-10", freq="M"))
    assert n_missing == 1
    row = mat[mat["B.CODE"] == "3402"].iloc[0]
    assert row["August - 26"] == 1 and row["October - 26"] == 1
    assert "September - 25" in labels


def test_full_report_each_category_has_its_own_month_range(tmp_path):
    """CONFIRMED by user (2026-10-08): ADR aur NET apni-apni ALAG continuous
    month-range dikhate hain -- jis mahine se us category ka pehla asli data
    shuru hota hai wahi se, doosri category ke start se match karne ke liye
    badhai nahi jaati."""
    import openpyxl
    from dbp import ingest

    path = str(tmp_path / "ranges.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PENDING BILLING"
    rows = [["S.NO", "CN_NO", "ADR DATE", "ADR BRANCH"], [1, "A1", "10-JUL-26", 101]]
    for r in rows:
        ws.append(r)
    ws2 = wb.create_sheet("POD ON NET BUT NOT BILLED")
    rows2 = [["S.NO", "CN_NO", "CN_DELIVERY_DATE", "BILL_BRANCH_CODE"], [1, "N1", "05-APR-26", 101]]
    for r in rows2:
        ws2.append(r)
    wb.save(path)

    data = ingest.load_workbook_categorized(path)
    report = ms.build_full_month_wise_report(data, {101: "X"},
                                               current_month_adr=pd.Period("2026-07", freq="M"),
                                               current_month_net=pd.Period("2026-07", freq="M"))
    adr_df, adr_months, _ = report["PENDING_BILLING"]
    net_df, net_months, _ = report["POD_NOT_BILLED"]
    assert adr_months == ["July - 26"]                 # ADR's own first/only month
    assert net_months[0] == "April - 26"                # NET starts from its own earliest
    assert "July - 26" in net_months                    # continuous through to current month


def test_full_report_party_rows_combined_with_branch_rows_and_zero_backfill(tmp_path):
    """Party rows (no S.NO/B.CODE) sit above numbered branch rows; a party
    present in one category but with zero in the other still gets a
    zero-filled row there (matching the reference report's display)."""
    import openpyxl
    from dbp import ingest

    path = str(tmp_path / "combined.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOSCH LTD"
    rows = [[101, "ADR MADE BUT NOT BILLED"] + [None] * 4, [None] * 6,
            ["S.NO", "CN_NO", "REMARK", "ADR DATE", "SOURCE_BRANCH", None],
            [1, "B1", "r", "10-JUL-26", "X", None]]
    for r in rows:
        ws.append(r)
    ws2 = wb.create_sheet("PENDING BILLING")
    rows2 = [["S.NO", "CN_NO", "ADR DATE", "ADR BRANCH"], [1, "BR1", "10-JUL-26", 101]]
    for r in rows2:
        ws2.append(r)
    wb.save(path)

    data = ingest.load_workbook_categorized(path)
    report = ms.build_full_month_wise_report(data, {101: "X"},
                                               current_month_adr=pd.Period("2026-07", freq="M"),
                                               current_month_net=pd.Period("2026-07", freq="M"))
    adr_df, _, _ = report["PENDING_BILLING"]
    net_df, _, _ = report["POD_NOT_BILLED"]

    assert list(adr_df["B.NAME"])[0] == "BOSCH LTD"    # party row first
    assert adr_df.iloc[0]["S.NO"] == "" and adr_df.iloc[0]["B.CODE"] == ""
    assert adr_df.iloc[1]["S.NO"] == 1                 # branch row numbered
    assert int(adr_df[adr_df["B.NAME"] == "TOTAL"].iloc[0]["TOTAL"]) == 2

    # BOSCH LTD has zero NET data -- still appears, zero-filled
    assert "BOSCH LTD" in list(net_df["B.NAME"])
    bosch_net = net_df[net_df["B.NAME"] == "BOSCH LTD"].iloc[0]
    assert int(bosch_net["TOTAL"]) == 0


def test_build_net_month_summary_end_to_end(tmp_path):
    import openpyxl
    from dbp import ingest

    path = str(tmp_path / "net_repro.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "POD ON NET BUT NOT BILLED"
    rows = [["S.NO", "CN_NO", "CN_DELIVERY_DATE", "BILL_BRANCH_CODE"]]
    for i in range(1, 4):
        rows.append([i, f"CN{i}", "10-AUG-26", 3402])
    for r in rows:
        ws.append(r)
    wb.save(path)

    categorized = ingest.load_workbook_categorized(path)
    mat, n_missing, labels = ms.build_net_month_summary(
        categorized, current_month=pd.Period("2026-10", freq="M"))
    assert n_missing == 0
    row = mat[mat["B.CODE"] == "3402"].iloc[0]
    assert row["TOTAL"] == 3
