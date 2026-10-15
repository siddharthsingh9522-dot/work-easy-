"""
CONFIRMED by user (2026-09-29): 'OTHER END' / 'DISPUTE MATTER' jaisa koi bhi
block/sub-section -- chahe apni poori alag sheet ho ya kisi PENDING
BILLING/POD/customer sheet ke BEECH mein ek chhota title-wala mini-section --
kabhi summary ke kisi category count mein nahi aana chahiye, aur na hi uski
rows kisi pichle (preceding) block mein chupke se jud kar us block ka count
badha deni chahiye.

Yeh file woh exact bug reproduce karti hai jo user ne report kiya tha: ek
sheet mein 119 asli PENDING BILLING rows the, lekin beech mein ek "DISPUTE
MATTER" sub-block (jise purana code block-marker hi nahi maanta tha) ki 5
rows pichle ADR block mein absorb ho kar summary mein 124 dikha rahi thi.
"""
import os
import openpyxl
import pandas as pd
import pytest

from dbp import ingest

HERE = os.path.dirname(os.path.abspath(__file__))


def _write_pending_billing_with_dispute(path, dispute_title="DISPUTE MATTER"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PENDING BILLING"
    rows = []
    rows.append([101, "ADR MADE BUT NOT BILLED"] + [None] * 5)
    rows.append([None] * 7)
    rows.append(["S.NO", "CN_NO", "CN_DATE", "BRANCH CODE", None, None, None])
    for i in range(1, 120):  # 119 genuine pending rows
        rows.append([i, f"CN{i:05d}", "01-JAN-26", 101, None, None, None])
    rows.append([None] * 7)
    rows.append([dispute_title] + [None] * 6)
    rows.append([None] * 7)
    rows.append(["S.NO", "CN_NO", "CN_DATE", "BRANCH CODE", None, None, None])
    for i in range(1, 6):  # 5 rows that must NOT be counted
        rows.append([i, f"DISP{i:03d}", "02-JAN-26", 101, None, None, None])
    for r in rows:
        ws.append(r)
    wb.save(path)


@pytest.mark.parametrize("dispute_title", [
    "DISPUTE MATTER", "Dispute Matter", "DISPUTE MATTERS", "DISPUTED MATTER",
    "DISP. MATTER", "DISP MATTER", "OTHER END", "OTHER-END", "O/E",
])
def test_dispute_block_excluded_from_pending_billing_count(tmp_path, dispute_title):
    path = str(tmp_path / "repro.xlsx")
    _write_pending_billing_with_dispute(path, dispute_title)
    data = ingest.load_workbook_categorized(path)

    pb = data.get("PENDING_BILLING")
    assert pb is not None
    # exact sheet-row count, not inflated by the dispute/other-end sub-block
    assert len(pb) == 119

    excluded = data.get("EXCLUDED_BLOCKS_DATA", {})
    assert "PENDING BILLING" in excluded
    assert len(excluded["PENDING BILLING"]) == 5


def test_dispute_block_inside_customer_sheet_is_split_per_block(tmp_path):
    """Ek customer-tab (jaise TATA MOTORS) ke andar ADR block + DISPUTE
    MATTER block dono ho to sirf dispute block excluded_blocks_data mein
    jaaye, ADR wala normal OTHER_BLOCKS_DATA mein hi rahe."""
    path = str(tmp_path / "repro_customer.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TATA MOTORS"
    rows = []
    rows.append(["ADR MADE BUT NOT BILLED"] + [None] * 6)
    rows.append([None] * 7)
    rows.append(["S.NO", "CN_NO", "CN_DATE", "BRANCH CODE", None, None, None])
    for i in range(1, 4):
        rows.append([i, f"TM{i:03d}", "01-JAN-26", 301, None, None, None])
    rows.append([None] * 7)
    rows.append(["DISPUTE MATTER"] + [None] * 6)
    rows.append([None] * 7)
    rows.append(["S.NO", "CN_NO", "CN_DATE", "BRANCH CODE", None, None, None])
    for i in range(1, 3):
        rows.append([i, f"TMD{i:03d}", "01-JAN-26", 301, None, None, None])
    for r in rows:
        ws.append(r)
    wb.save(path)

    data = ingest.load_workbook_categorized(path)
    other = data.get("OTHER_BLOCKS_DATA", {}).get("TATA MOTORS")
    excluded = data.get("EXCLUDED_BLOCKS_DATA", {}).get("TATA MOTORS")
    assert other is not None and len(other) == 3
    assert excluded is not None and len(excluded) == 2
    assert set(other["_BLOCK_TITLE"].unique()) == {"ADR MADE BUT NOT BILLED"}
    assert set(excluded["_BLOCK_TITLE"].unique()) == {"DISPUTE MATTER"}


def test_party_sheet_adr_and_net_blocks_both_count_under_party(tmp_path):
    """Image-shaped repro (2026-10-04): a party sheet (e.g. 'TATA MOTOR')
    with branch marker 2601 holding BOTH an 'ADR MADE BUT NOT BILLED' block
    and a 'POD ON NET BUT NOT BILLED' block -- each must land in its own
    summary category under the PARTY row, never the branch."""
    import openpyxl
    from dbp import ingest
    from dbp import report_builder

    path = str(tmp_path / "party_repro.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TATA MOTOR"
    rows = [[2601, "ADR MADE BUT NOT BILLED"] + [None] * 4, [None] * 6,
            ["S.NO", "CN_NO", "REMARK", "CN_DATE", "SOURCE_BRANCH", None]]
    for i in range(1, 6):
        rows.append([i, f"ADRCN{i}", "remark", "11.04.2026", "ZIRAKPUR I-PUNJAB", None])
    rows += [[None] * 6, [2601, "POD ON NET BUT NOT BILLED"] + [None] * 4, [None] * 6,
             ["S.NO", "CN_NO", "REMARK", "CN_DATE", "SOURCE_BRANCH", None]]
    for i in range(1, 4):
        rows.append([i, f"NETCN{i}", "remark", "13.04.2026", "BANGLORE HUB-KARNATAKA", None])
    for r in rows:
        ws.append(r)

    ws2 = wb.create_sheet("PENDING BILLING")
    rows2 = [["S.NO", "CN_NO", "CN_DATE", "ADR BRANCH"]]
    for i in range(1, 4):
        rows2.append([i, f"BRCN{i}", "01.05.2026", 2601])
    for r in rows2:
        ws2.append(r)
    wb.save(path)

    data = ingest.load_workbook_categorized(path)
    out = report_builder.build_summary(data, {2601: "ZIRAKPUR"}, {2601: "ZIRAKPUR"})

    branch_row = out[pd.to_numeric(out["B.CODE"], errors="coerce") == 2601].iloc[0]
    assert int(branch_row["PENDING BILLING"]) == 3   # only the branch's own direct sheet

    party_row = out[out["B.NAME"].astype(str).str.upper() == "TATA MOTOR"].iloc[0]
    assert int(party_row["PENDING BILLING"]) == 5
    assert int(party_row["POD ON NET BUT NOT BILLED"]) == 3


@pytest.mark.parametrize("title", ["OTHER", "DISPUTE"])
def test_bare_other_and_dispute_short_forms_excluded(tmp_path, title):
    """CONFIRMED by user (2026-10-07): block kabhi sirf 'OTHER' ya sirf
    'DISPUTE' (bina 'END'/'MATTER' suffix ke) likha hota hai -- yeh bhi
    pehchaana jaana chahiye."""
    path = str(tmp_path / f"bare_{title}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PENDING BILLING"
    rows = [[101, "ADR MADE BUT NOT BILLED"] + [None] * 5, [None] * 7,
            ["S.NO", "CN_NO", "CN_DATE", "ADR BRANCH", None, None, None]]
    for i in range(1, 6):
        rows.append([i, f"CN{i}", "01-JAN-26", 101, None, None, None])
    rows += [[None] * 7, [title] + [None] * 6, [None] * 7,
             ["S.NO", "CN_NO", "CN_DATE", "ADR BRANCH", None, None, None]]
    for i in range(1, 3):
        rows.append([i, f"X{i}", "02-JAN-26", 101, None, None, None])
    for r in rows:
        ws.append(r)
    wb.save(path)

    from dbp import ingest
    data = ingest.load_workbook_categorized(path)
    assert len(data["PENDING_BILLING"]) == 5
    assert len(data.get("EXCLUDED_BLOCKS_DATA", {}).get("PENDING BILLING", [])) == 2


def test_party_short_form_names_merge_into_one_summary_row_and_tab(tmp_path):
    """CONFIRMED by user (2026-10-07): 'TATA MOTOR' / 'TATA MOTORS' / 'TATA'
    sab ek hi party hain -- ek hi SUMMARY row, ek hi export tab, aur export
    tab ka row-count summary ke TOTAL se exactly match karna chahiye."""
    import openpyxl as oxl
    from dbp import ingest
    from dbp import report_builder

    def mk(path, sheet_name, title, n, prefix):
        wb = oxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        rows = [[101, title] + [None] * 4, [None] * 6,
                ["S.NO", "CN_NO", "REMARK", "CN_DATE", "SOURCE_BRANCH", None]]
        for i in range(1, n + 1):
            rows.append([i, f"{prefix}{i}", "r", "11.04.2026", "ZIRAKPUR", None])
        for r in rows:
            ws.append(r)
        wb.save(path)

    p1 = str(tmp_path / "f1.xlsx")
    p2 = str(tmp_path / "f2.xlsx")
    mk(p1, "TATA MOTOR", "ADR MADE BUT NOT BILLED", 5, "ADR")
    mk(p2, "TATA", "POD ON NET BUT NOT BILLED", 3, "NET")

    data = ingest.merge_all_uploaded([p1, p2])
    summary = report_builder.build_summary(data, {}, {})

    party_rows = summary[summary["B.NAME"].astype(str).str.upper().str.startswith("TATA")]
    assert len(party_rows) == 1   # not two separate rows for "TATA MOTOR" and "TATA"
    assert int(party_rows.iloc[0]["TOTAL"]) == 8

    out_path = str(tmp_path / "final.xlsx")
    report_builder.export_full_report_to_excel(data, summary, out_path)
    wb = oxl.load_workbook(out_path)
    assert "TATA MOTOR" in wb.sheetnames and "TATA" not in wb.sheetnames
    n_rows = sum(1 for r in wb["TATA MOTOR"].iter_rows(min_row=2, values_only=True)
                 if any(c is not None for c in r))
    assert n_rows == 8   # exactly matches the summary TOTAL -- no mismatch


def test_summary_never_empty_when_only_party_sheets_exist(tmp_path):
    """Real bug found 2026-10-07: build_summary() used to return a totally
    empty frame whenever there were zero direct branch-level sheets, even if
    party/customer data existed -- silently dropping it. A workbook that is
    100% party tabs (no PENDING BILLING/POD/etc sheet at all) must still
    produce a SUMMARY with that party's row."""
    import openpyxl as oxl
    from dbp import ingest
    from dbp import report_builder

    path = str(tmp_path / "only_party.xlsx")
    wb = oxl.Workbook()
    ws = wb.active
    ws.title = "BOSCH"
    rows = [[201, "ADR MADE BUT NOT BILLED"] + [None] * 4, [None] * 6,
            ["S.NO", "CN_NO", "REMARK", "CN_DATE", "SOURCE_BRANCH", None]]
    for i in range(1, 4):
        rows.append([i, f"B{i}", "r", "01.01.2026", "X", None])
    for r in rows:
        ws.append(r)
    wb.save(path)

    data = ingest.load_workbook_categorized(path)
    summary = report_builder.build_summary(data, {}, {})
    assert not summary.empty
    bosch = summary[summary["B.NAME"].astype(str).str.upper() == "BOSCH"]
    assert len(bosch) == 1
    assert int(bosch.iloc[0]["PENDING BILLING"]) == 3


def test_dispatch_sheet_not_falsely_excluded():
    """Safety check: 'disp' ko akela keyword kabhi nahi rakha (sirf 'disp
    matter'/'disp. matter') taaki ek 'DISPATCH' jaisi asli sheet galti se
    EXCLUDED na ban jaaye."""
    from dbp import ingest as ing
    assert ing.categorize_sheet("DISPATCH REPORT") != "EXCLUDED"
    assert ing._match_block_title_category("DISPATCH PENDING") != "EXCLUDED"
