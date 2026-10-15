"""DBP tools end-to-end through the web layer (upload -> report -> validation -> chatbot -> month-wise -> masters)."""
import io
import re

import openpyxl
import pytest

from conftest import admin_post, do_login, do_signup
from core.store import user_id_for
from dbp import masters


def make_xlsx(dateless=False):
    """A small but realistic export: one PENDING BILLING block for branch 101 (+ optionally one date-less docket)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PENDING BILLING"
    ws.append([101, "ADR MADE BUT NOT BILLED"] + [None] * 5)
    ws.append([None] * 7)
    ws.append(["S.NO", "CN_NO", "ADR DATE", "ADR BRANCH", None, None, None])
    ws.append([1, "CN1", "10-AUG-26", 101, None, None, None])
    ws.append([2, "CN2", "12-SEP-26", 101, None, None, None])
    ws.append([3, "CN3", None if dateless else "01-OCT-26", 101, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def token(client, path):
    html = client.get(path).get_data(as_text=True)
    return re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)


def upload(client, url, page, fields=None, files=(("x.xlsx", None),), dateless=False, name="files"):
    data = dict(csrf_token=token(client, page), **(fields or {}))
    data[name] = [(io.BytesIO(content if content is not None else make_xlsx(dateless)), fn) for fn, content in files]
    return client.post(url, data=data, content_type="multipart/form-data")


def post(client, url, page, **fields):
    return client.post(url, data=dict(csrf_token=token(client, page), **fields))


@pytest.fixture()
def made(admin):
    r = upload(admin, "/dbp/report/generate", "/dbp/report", {"region": "Test Region"})
    assert r.status_code == 302
    return admin


def test_report_generate_page_and_download(made):
    page = made.get("/dbp/report").get_data(as_text=True)
    assert "Test Region" in page and "101" in page and "SUMMARY preview" in page
    dl = made.get("/dbp/report/download")
    assert dl.status_code == 200 and dl.data[:2] == b"PK"
    names = openpyxl.load_workbook(io.BytesIO(dl.data)).sheetnames
    assert "SUMMARY" in [n.upper() for n in names]


def test_non_excel_upload_is_rejected(admin):
    r = upload(admin, "/dbp/report/generate", "/dbp/report", {"region": "R"}, files=(("notes.txt", b"hello"),))
    assert r.status_code == 302
    assert "Excel" in admin.get("/dbp/report").get_data(as_text=True)


def test_validation_runs_and_page_renders(made):
    assert post(made, "/dbp/validation", "/dbp/validation").status_code == 302
    page = made.get("/dbp/validation").get_data(as_text=True)
    assert "Validation poora hua" in page or "issue" in page.lower()


def test_chatbot_answers_from_data(made):
    tok = token(made, "/dbp/chatbot")
    r = made.post("/dbp/chatbot/ask", json={"q": "101 ki pending batao"}, headers={"X-CSRF-Token": tok})
    assert r.status_code == 200 and "101" in r.get_json()["answer"]
    assert "101 ki pending" in made.get("/dbp/chatbot").get_data(as_text=True)      # history persisted


def test_chatbot_needs_a_report_first(admin):
    r = admin.post("/dbp/chatbot/ask", json={"q": "hi"}, headers={"X-CSRF-Token": token(admin, "/dbp/chatbot")})
    assert r.status_code == 400


def test_month_wise_asks_for_month_then_builds(admin):
    r = upload(admin, "/dbp/month/run", "/dbp/month?cat=adr", {"cat": "adr", "region": "Test Region"}, dateless=True)
    assert r.status_code == 302
    page = admin.get("/dbp/month?cat=adr").get_data(as_text=True)
    assert "kis month tak" in page                                  # date-less docket -> asks the month
    post(admin, "/dbp/month/period", "/dbp/month?cat=adr", cat="adr", month_single="10", year_single="2026")
    page = admin.get("/dbp/month?cat=adr").get_data(as_text=True)
    assert "kis month tak" not in page and "TOTAL" in page and "August - 26" in page
    dl = admin.get("/dbp/month/download?cat=adr")
    assert dl.status_code == 200 and dl.data[:2] == b"PK"


def test_month_wise_without_dateless_rows_goes_straight_to_result(admin):
    upload(admin, "/dbp/month/run", "/dbp/month?cat=adr", {"cat": "adr", "region": "R"})
    page = admin.get("/dbp/month?cat=adr").get_data(as_text=True)
    assert "kis month tak" not in page and "TOTAL" in page


def test_month_wise_full_region_renders_with_audit(admin):
    upload(admin, "/dbp/month/run", "/dbp/month?cat=full", {"cat": "full", "region": "R"})
    page = admin.get("/dbp/month?cat=full")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "kis month tak" in html or "Counting audit" in html


def test_each_user_has_their_own_workspace(app, made):
    other = app.test_client()
    do_signup(other, "Ravi", "ravi@x.com")
    admin_post(made, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve",
               modules=["dbp_report", "ai_chatbot"])
    do_login(other, "ravi@x.com")
    assert "Abhi koi report generate nahi hui" in other.get("/dbp/report").get_data(as_text=True)


def test_dbp_tools_respect_module_access(app, admin):
    u = app.test_client()
    do_signup(u, "Ravi", "ravi@x.com")
    admin_post(admin, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve", modules="paid_topay")
    do_login(u, "ravi@x.com")
    for path in ("/dbp/report", "/dbp/validation", "/dbp/month", "/dbp/chatbot", "/dbp/directory"):
        assert u.get(path).status_code == 403, path


def test_directory_is_admin_only_and_bad_upload_keeps_old_file(app, admin):
    u = app.test_client()
    do_signup(u, "Ravi", "ravi@x.com")
    admin_post(admin, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve",
               modules=["dbp_report", "month_summary", "ai_chatbot", "validation", "directory"])
    do_login(u, "ravi@x.com")
    assert u.get("/dbp/directory").status_code == 403               # admin_only even if ticked
    before = len(masters.load().directory_df)
    assert before > 1000
    r = upload(admin, "/dbp/directory", "/dbp/directory", {"which": "directory"},
               files=(("junk.xlsx", b"this is not an excel file"),), name="file")
    assert r.status_code == 302
    assert "purani file safe" in admin.get("/dbp/directory").get_data(as_text=True)
    assert len(masters.load().directory_df) == before


def test_directory_search(admin):
    page = admin.get("/dbp/directory?q=zzzz-no-such-person").get_data(as_text=True)
    assert "Directory me kuch nahi mila" in page
    assert "Directory rows" in admin.get("/dbp/directory").get_data(as_text=True)
