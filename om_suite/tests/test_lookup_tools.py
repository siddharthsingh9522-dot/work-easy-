"""Branch Finder, Customer Search and OM Automation against a fake internal portal (real HTTP)."""
import io
import json
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, HTTPServer
from urllib.parse import parse_qs

import openpyxl
import pytest

from app import create_app
from blueprints import branch as branch_mod
from blueprints import customer as customer_mod
from conftest import admin_post, do_login, do_signup
from core import jobs
from core.store import LocalStore, user_id_for
from omauto.api import gst_captcha

BRANCHES = {"1001": ("Delhi HO", "Addr 1", "1001", "07AAA"), "1002": ("Pune", "Addr 2", "1001", "27BBB")}
# 22 customer columns, in the portal's fixed order
CUSTOMERS = {"C1": ["C1", "Acme Ltd", "Some road", "BK1", "AC1", "BL1", "Yes", "Regular", "07AAACA1234A1Z5", "Active",
                    "M1", "K", "Doc", "01-01-2024", "V1", "99999", "MG", "DC", "RT", "5", "user", "02-02-2024"]}


def table(headers, rows, css=""):
    th = "".join(f"<th>{h}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
    return f'<html><body><table {css}><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table></body></html>'


class Portal:
    def __init__(self):
        outer = self
        self.hits = []

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _send(self, body, ctype="text/html", code=200):
                data = body.encode()
                self.send_response(code)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def do_POST(self):
                form = {k: v[0] for k, v in parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode(), keep_blank_values=True).items()}
                outer.hits.append((self.path, form))
                if self.path.startswith("/branch"):
                    b = BRANCHES.get(form.get("bcode"))
                    rows = [["1", form["bcode"], *b[:2], b[2], b[3]]] if b else []
                    return self._send(table(["#", "code", "name", "addr", "ctrl", "gst"], rows, 'class="table"'))
                if self.path.startswith("/customer"):
                    code = form.get("ccode") or form.get("blcode")
                    c = CUSTOMERS.get(code)
                    if not c:
                        return self._send("<html>No Data Found</html>")
                    return self._send(table(["h"] * 22, [c]))
                self._send("", code=404)

            def do_GET(self):
                m = re.match(r"/cn/(\w+)$", self.path)
                if not m:
                    return self._send("", code=404)
                if m.group(1) == "CN404":
                    return self._send(json.dumps({"status": "error", "message": "CN Not Found"}), "application/json")
                main = {"cnDate": "01-01-2026", "billingParty": "C1-Acme Ltd", "source": "101-Delhi", "destination": "202-Pune",
                        "consignor": "X1-Sender", "consignee": "Y1-Receiver", "currentStatus": "Delivered"}
                return self._send(json.dumps({"status": "success", "data": {"main": main, "freight": {}, "billFreight": []}}),
                                  "application/json")

        self.server = HTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.server.server_port}"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()


@pytest.fixture()
def portal():
    p = Portal()
    yield p
    p.server.shutdown()


@pytest.fixture(autouse=True)
def fast(monkeypatch):
    monkeypatch.setattr(branch_mod, "DELAY_RANGE", (0, 0))
    monkeypatch.setattr(branch_mod, "RETRY_COUNT", 1)
    monkeypatch.setattr(customer_mod, "DELAY_SECONDS", 0)
    monkeypatch.setattr(customer_mod, "NOT_FOUND_WAIT", 0)
    gst_captcha.SESSIONS.clear()


@pytest.fixture()
def lapp(tmp_path, portal):
    return create_app({"TESTING": True, "USER_CACHE_TTL": 0, "ADMIN_EMAILS": {"admin@x.com"}, "SECRET_KEY": "t",
                       "DATA_DIR": tmp_path / "data", "BRANCH_URL": portal.url + "/branch",
                       "CUSTOMER_URL": portal.url + "/customer", "CN_API_URL": portal.url + "/cn/{}"},
                      store=LocalStore(tmp_path / "t.db"))


@pytest.fixture()
def me(lapp):
    c = lapp.test_client()
    do_signup(c, "Admin", "admin@x.com")
    do_login(c, "admin@x.com")
    return c


def tok(c):
    return re.search(r'csrf-token" content="([^"]+)"', c.get("/").get_data(as_text=True)).group(1)


def xlsx(rows, header=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    if header:
        ws.append(header)
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return b


def post_file(c, url, rows, header=None, name="in.xlsx"):
    return c.post(url, data={"file": (xlsx(rows, header), name)}, content_type="multipart/form-data",
                  headers={"X-CSRF-Token": tok(c)})


def wait(c, job_id, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        s = c.get(f"/jobs/{job_id}/status").get_json()
        if s["status"] != "running":
            return s
        time.sleep(0.1)
    raise AssertionError("job did not finish")


# ================================================================= Branch Finder
def test_branch_single_lookup(me):
    r = me.post("/branch/lookup", data={"branch_code": "1001"}, headers={"X-CSRF-Token": tok(me)}).get_json()
    assert r["ok"] and r["data"]["BRANCH_NAME"] == "Delhi HO" and r["data"]["GST_NO"] == "07AAA"
    assert not me.post("/branch/lookup", data={"branch_code": "9999"}, headers={"X-CSRF-Token": tok(me)}).get_json()["ok"]
    assert "numeric" in me.post("/branch/lookup", data={"branch_code": "abc"}, headers={"X-CSRF-Token": tok(me)}).get_json()["error"]


def test_branch_bulk_job_and_report(me):
    r = post_file(me, "/branch/bulk/start", [[1001], [1002], [9999]], header=["Branch Code"]).get_json()
    s = wait(me, r["job_id"])
    assert s["status"] == "done" and s["extra"] == {"success": 2, "errors": 1} and s["total"] == 3
    wb = openpyxl.load_workbook(io.BytesIO(me.get(s["download_url"]).data))
    assert wb.sheetnames == ["Success Report", "Error Report", "Summary"]
    assert wb["Success Report"].max_row == 3 and wb["Error Report"].max_row == 2


def test_branch_rejects_non_excel(me):
    r = me.post("/branch/bulk/start", data={"file": (io.BytesIO(b"x"), "a.txt")}, content_type="multipart/form-data",
                headers={"X-CSRF-Token": tok(me)})
    assert r.status_code == 400


def test_branch_download_single_result(me):
    r = me.post("/branch/lookup/download", data={"csrf_token": tok(me), "SEARCH_CODE": "1001", "BRANCH_NAME": "Delhi HO"})
    assert r.status_code == 200 and r.data[:2] == b"PK"


def test_jobs_are_private_to_their_owner(lapp, me):
    s = wait(me, post_file(me, "/branch/bulk/start", [[1001]], header=["c"]).get_json()["job_id"])
    other = lapp.test_client()
    do_signup(other, "Ravi", "ravi@x.com")
    admin_post(me, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve", modules=["branch_finder"])
    do_login(other, "ravi@x.com")
    jid = s["job_id"]
    assert other.get(f"/jobs/{jid}/status").status_code == 404
    assert other.get(f"/jobs/{jid}/download").status_code == 404
    assert other.post(f"/jobs/{jid}/stop", headers={"X-CSRF-Token": tok(other)}).status_code == 404


# ================================================================= Customer Search
def test_customer_manual_search(me):
    ok = me.post("/customer/api/manual-search", json={"ccode": "C1"}, headers={"X-CSRF-Token": tok(me)}).get_json()
    assert ok["ok"] and ok["rows"][0]["CNAME"] == "Acme Ltd" and ok["rows"][0]["GST NO"].startswith("07")
    none = me.post("/customer/api/manual-search", json={"ccode": "ZZ"}, headers={"X-CSRF-Token": tok(me)}).get_json()
    assert none["ok"] and none["rows"] == []
    assert me.post("/customer/api/manual-search", json={}, headers={"X-CSRF-Token": tok(me)}).status_code == 400


def test_customer_portal_down_is_a_clean_502(lapp, me):
    lapp.config["CUSTOMER_URL"] = "http://127.0.0.1:9/customer"
    r = me.post("/customer/api/manual-search", json={"ccode": "C1"}, headers={"X-CSRF-Token": tok(me)})
    assert r.status_code == 502 and "jawab nahi" in r.get_json()["error"]


def test_customer_bulk_with_not_found_sheet(me):
    s = wait(me, post_file(me, "/customer/api/bulk-search/start", [["C1"], ["NOPE"]], header=["CCode"]).get_json()["job_id"])
    assert s["status"] == "done" and s["extra"] == {"found": 1, "not_found": 1}
    wb = openpyxl.load_workbook(io.BytesIO(me.get(s["download_url"]).data))
    assert wb.sheetnames == ["Results", "Not Found"] and wb["Not Found"]["A2"].value == "NOPE"
    assert wb["Results"]["C2"].value == "Acme Ltd"


def test_customer_bulk_without_code_column_fails_cleanly(me):
    s = wait(me, post_file(me, "/customer/api/bulk-search/start", [["abc", "def"]]).get_json()["job_id"])
    assert s["status"] == "error" and s["error"]


# ================================================================= OM Automation
def api(c, url, payload):
    return c.post(url, json=payload, headers={"X-CSRF-Token": tok(c)})


def test_om_manual_cn_party_and_master(me):
    cn = api(me, "/om/api/manual/cn", {"value": "CN123"}).get_json()
    assert cn["success"] and cn["data"]["CN"] == "CN123" and cn["data"]["Party Code"] == "C1"
    bad = api(me, "/om/api/manual/cn", {"value": "CN404"}).get_json()
    assert not bad["success"] and "CN Not Found" in bad["data"]["Error"]
    party = api(me, "/om/api/manual/party", {"value": "C1"}).get_json()
    assert party["success"] and party["data"]["Party Name"] == "Acme Ltd"
    assert api(me, "/om/api/manual/cn", {"value": ""}).status_code == 400
    assert api(me, "/om/api/manual/nope", {"value": "x"}).status_code == 404


def test_om_excel_batch_makes_a_report(me):
    r = post_file(me, "/om/api/excel/cn/start", [["CN001"], ["CN404"], ["CN003"]], header=["CN No"]).get_json()
    assert r["success"]
    s = wait(me, r["job_id"])
    assert s["status"] == "done" and s["current"] == 3 and any("Done. Success: 2 | Failed: 1" in l for l in s["log"])
    wb = openpyxl.load_workbook(io.BytesIO(me.get(s["download_url"]).data))
    assert "Success Report" in wb.sheetnames and wb["Success Report"].max_row == 3
    assert me.get(f"/om/api/excel/status/{r['job_id']}").get_json()["status"] == "done"


def test_om_excel_rejects_bad_files(me):
    r = me.post("/om/api/excel/cn/start", data={"file": (io.BytesIO(b"x"), "a.csv")}, content_type="multipart/form-data",
                headers={"X-CSRF-Token": tok(me)})
    assert r.status_code == 400 and not r.get_json()["success"]
    assert api(me, "/om/api/excel/zzz/start", {}).status_code == 404


# ---- GST captcha flow with a fake browser (no Chrome needed)
class FakeBrowser:
    def __init__(self):
        self.wrong = True
        self.closed = []
        self.page = ""


@pytest.fixture()
def gst(monkeypatch):
    fb = FakeBrowser()
    monkeypatch.setattr(gst_captcha, "open_browser", lambda: object())
    monkeypatch.setattr(gst_captcha, "go_to_search_page", lambda d: None)
    monkeypatch.setattr(gst_captcha, "fill_gst", lambda d, g: setattr(fb, "gst", g))
    monkeypatch.setattr(gst_captcha, "screenshot_captcha", lambda d, p: open(p, "wb").write(b"\x89PNG-fake"))
    monkeypatch.setattr(gst_captcha, "fill_captcha", lambda d, t: setattr(fb, "typed", t))
    monkeypatch.setattr(gst_captcha, "click_search", lambda d: None)
    monkeypatch.setattr(gst_captcha, "get_page_text", lambda d: "wrong" if fb.typed == "bad" else "ok")
    monkeypatch.setattr(gst_captcha, "looks_like_wrong_captcha", lambda t: t == "wrong")
    monkeypatch.setattr(gst_captcha, "extract_result",
                        lambda d, g: {"GSTIN Searched": g, "Trade Name": "ACME", "Remarks": "Active"})
    monkeypatch.setattr(gst_captcha, "close_browser", lambda d: fb.closed.append(d))
    return fb


def test_gst_captcha_flow_retry_then_done_and_report(me, gst):
    start = api(me, "/om/api/gst/captcha/start_manual", {"gstin": "07aaaca1234a1z5"}).get_json()
    assert start["success"] and start["status"] == "waiting_captcha" and start["gst"] == "07AAACA1234A1Z5"
    assert me.get("/om/api/gst/captcha/image").data.startswith(b"\x89PNG")
    retry = api(me, "/om/api/gst/captcha/submit", {"captcha": "bad"}).get_json()
    assert retry["status"] == "retry_captcha"
    done = api(me, "/om/api/gst/captcha/submit", {"captcha": "good"}).get_json()
    assert done["status"] == "done" and done["last_result"]["remarks"] == "Active"
    wb = openpyxl.load_workbook(io.BytesIO(me.get("/om/api/gst/report").data))
    assert wb.active.max_row == 2 and gst.closed                      # browser is closed when finished


def test_gst_validates_input_and_needs_a_session(me, gst):
    assert api(me, "/om/api/gst/captcha/start_manual", {"gstin": "short"}).status_code == 400
    assert api(me, "/om/api/gst/captcha/submit", {"captcha": "x"}).status_code == 400
    assert me.get("/om/api/gst/captcha/image").status_code == 404
    assert me.get("/om/api/gst/report").status_code == 404


def test_gst_browser_slots_are_limited_and_isolated_per_user(lapp, me, gst, monkeypatch):
    monkeypatch.setattr(gst_captcha, "MAX_SESSIONS", 1)
    assert api(me, "/om/api/gst/captcha/start_manual", {"gstin": "07AAACA1234A1Z5"}).get_json()["success"]
    other = lapp.test_client()
    do_signup(other, "Ravi", "ravi@x.com")
    admin_post(me, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve", modules=["om_automation"])
    do_login(other, "ravi@x.com")
    busy = api(other, "/om/api/gst/captcha/start_manual", {"gstin": "27BBBCB1234B1Z5"})
    assert busy.status_code == 503 and "doosre user" in busy.get_json()["message"]
    assert other.get("/om/api/gst/captcha/image").status_code == 404           # cannot see my captcha
    api(me, "/om/api/gst/captcha/stop", {})                                     # I finish -> slot is free again
    assert api(other, "/om/api/gst/captcha/start_manual", {"gstin": "27BBBCB1234B1Z5"}).get_json()["success"]


def test_gst_idle_sessions_are_reaped(me, gst, monkeypatch):
    api(me, "/om/api/gst/captcha/start_manual", {"gstin": "07AAACA1234A1Z5"})
    for st in gst_captcha.SESSIONS.values():
        st["last"] -= 10_000
    gst_captcha.cleanup_stale()
    assert not gst_captcha.SESSIONS and gst.closed


def test_gst_without_chrome_gives_a_clear_error(me, monkeypatch):
    def boom():
        raise RuntimeError("GST captcha search needs Selenium + Chrome")
    monkeypatch.setattr(gst_captcha, "open_browser", boom)
    r = api(me, "/om/api/gst/captcha/start_manual", {"gstin": "07AAACA1234A1Z5"})
    assert r.status_code == 503 and "Chrome" in r.get_json()["message"]


# ================================================================= access control + pages
def test_new_tools_require_their_module(lapp, me):
    u = lapp.test_client()
    do_signup(u, "Ravi", "ravi@x.com")
    admin_post(me, f"/admin/users/{user_id_for('ravi@x.com')}/action", action="approve", modules=["paid_topay"])
    do_login(u, "ravi@x.com")
    for path in ("/branch/", "/customer/", "/om/"):
        assert u.get(path).status_code == 403, path
    assert u.post("/om/api/manual/cn", json={"value": "x"}, headers={"X-CSRF-Token": tok(u)}).status_code == 403


def test_pages_render_and_sidebar_lists_om_tools(me):
    for path in ("/branch/", "/customer/", "/om/"):
        assert me.get(path).status_code == 200, path
    html = me.get("/om/").get_data(as_text=True)
    assert 'data-sub-mode="gst"' in html and 'id="pages-data"' in html


# ================================================================= regression: "Unexpected token '<'" bug
def test_api_calls_without_csrf_get_json_not_an_html_page(me):
    """A stale cached script that forgets the CSRF header used to receive an HTML 400 page -> 'Unexpected token <'."""
    for url, body in (("/om/api/manual/cn", {"value": "CN001"}), ("/om/api/manual/party", {"value": "C1"}),
                      ("/customer/api/manual-search", {"ccode": "C1"}), ("/dbp/chatbot/ask", {"q": "hi"})):
        r = me.post(url, json=body)
        assert r.status_code == 400 and r.mimetype == "application/json", url
        assert r.get_json()["success"] is False and "refresh" in r.get_json()["message"]


def test_api_calls_when_logged_out_get_json_401_not_a_login_redirect(lapp):
    anon = lapp.test_client()
    t = re.search(r'name="csrf_token" value="([^"]+)"', anon.get("/login").get_data(as_text=True)).group(1)
    for url in ("/om/api/manual/cn", "/customer/api/manual-search", "/branch/lookup"):
        r = anon.post(url, json={"value": "x"}, headers={"X-CSRF-Token": t})
        assert r.status_code == 401 and r.mimetype == "application/json" and "login" in r.get_json()["message"].lower(), url
    assert anon.get("/jobs/abc/status").status_code == 401 and anon.get("/jobs/abc/status").mimetype == "application/json"
    assert anon.get("/om/").status_code == 302                       # normal pages still redirect to the login page


def test_unknown_api_url_and_server_errors_are_json_too(me, monkeypatch):
    assert me.get("/om/api/nope").mimetype == "application/json"
    monkeypatch.setitem(__import__("omauto.core.batch_runner", fromlist=["x"]).MANUAL_FETCHERS, "cn",
                        lambda v: (_ for _ in ()).throw(RuntimeError("boom")))
    r = api(me, "/om/api/manual/cn", {"value": "CN001"})
    assert r.status_code == 500 and r.mimetype == "application/json" and "boom" in r.get_json()["message"]


def test_static_urls_change_when_a_file_changes(lapp, me):
    import pathlib
    html = me.get("/").get_data(as_text=True)
    v1 = re.search(r"app\.js\?v=([0-9a-f]+)", html).group(1)
    assert v1 != "1.0.0-phase1" and len(v1) == 10                    # content hash, not a constant
    f = pathlib.Path(lapp.static_folder) / "js" / "app.js"
    original = f.read_bytes()
    try:
        f.write_bytes(original + b"\n/* edited */\n")
        v2 = re.search(r"app\.js\?v=([0-9a-f]+)", me.get("/").get_data(as_text=True)).group(1)
        assert v2 != v1                                              # a new URL -> stale cached copies are bypassed
    finally:
        f.write_bytes(original)


def test_service_worker_cache_was_rotated(me):
    assert b"omsuite-static-v2" in me.get("/sw.js").data


def test_party_lookup_sends_what_the_browser_sends(me, portal):
    r = api(me, "/om/api/manual/party", {"value": "C1"}).get_json()
    assert r["success"]
    path, form = [h for h in portal.hits if h[0].startswith("/customer")][-1]
    assert form == {"cname": "", "ccode": "C1", "lcode": "", "blcode": "", "mcode": "", "custgst": "", "submit": "Search"}


def test_party_lookup_headers_are_the_portals_not_the_cn_apis(me, portal):
    seen = {}
    import socket
    s = socket.socket(); s.bind(("127.0.0.1", 0)); port = s.getsockname()[1]; s.close()

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a): pass
        def do_POST(self):
            self.rfile.read(int(self.headers.get("Content-Length", 0)))
            seen.update({k.lower(): v for k, v in self.headers.items()})
            body = b"<html><title>Login</title>please sign in</html>"
            self.send_response(200); self.send_header("Content-Type", "text/html"); self.send_header("Content-Length", str(len(body)))
            self.end_headers(); self.wfile.write(body)
    srv = HTTPServer(("127.0.0.1", port), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        import omauto.config as oc
        oc.CUSTOMER_API = f"http://127.0.0.1:{port}/oracle/query/customer.php?u=39942"
        r = api(me, "/om/api/manual/party", {"value": "734032"}).get_json()
    finally:
        srv.shutdown()
    assert seen["accept"].startswith("text/html") and seen["referer"] == oc.CUSTOMER_API
    assert seen["origin"] == f"http://127.0.0.1:{port}" and "omone" not in seen.get("origin", "")
    assert seen["content-type"] == "application/x-www-form-urlencoded"
    # a login/error page is reported as such instead of the misleading "Customer Not Found"
    assert not r["success"] and "expected result page" in r["data"]["Error"] and "Login" in r["data"]["Error"]
